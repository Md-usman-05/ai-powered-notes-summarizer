from io import BytesIO
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader


class UnsupportedFileError(ValueError):
    pass


def extract_text_from_upload(uploaded_file):
    """
    Extract readable text from an uploaded document.

    Supported:
    - PDF
    - DOCX
    - HTML
    - TXT
    - CSV
    - JSON
    - Markdown
    - source/code files
    """

    if not uploaded_file:
        raise UnsupportedFileError(
            "No file was uploaded."
        )

    file_name = uploaded_file.name or "uploaded-file"
    suffix = Path(file_name).suffix.lower()

    content = uploaded_file.read()

    if not content:
        raise UnsupportedFileError(
            "The uploaded file is empty."
        )

    if suffix == ".pdf":
        return extract_pdf_text(content)

    if suffix == ".docx":
        return extract_docx_text(content)

    if suffix in {".html", ".htm"}:
        return extract_html_text(content)

    return extract_plain_text(content)


# ============================================================
# PDF
# ============================================================

def extract_pdf_text(content):
    try:
        reader = PdfReader(BytesIO(content))
    except Exception as exc:
        raise UnsupportedFileError(
            "The PDF could not be opened. Please upload a valid PDF."
        ) from exc

    if not reader.pages:
        raise UnsupportedFileError(
            "The PDF does not contain any pages."
        )

    pages = []

    for page_number, page in enumerate(reader.pages, start=1):

        try:
            # Layout mode generally gives better results for
            # educational notes and structured documents.
            page_text = page.extract_text(
                extraction_mode="layout"
            ) or ""

        except TypeError:
            # Compatibility fallback for pypdf versions
            # that don't support extraction_mode.
            page_text = page.extract_text() or ""

        except Exception:
            page_text = ""

        page_text = page_text.strip()

        if page_text:
            pages.append(
                f"[Page {page_number}]\n{page_text}"
            )

    text = "\n\n".join(pages).strip()

    # Check whether meaningful text was extracted.
    if not text:
        raise UnsupportedFileError(
            "This PDF appears to be scanned or image-based and does not contain selectable text. "
            "Please upload a text-based PDF or paste the notes manually."
        )

    # Avoid accepting PDFs where only a few random characters
    # were extracted.
    words = text.split()

    if len(words) < 10:
        raise UnsupportedFileError(
            "Very little readable text could be extracted from this PDF. "
            "It may be a scanned/image-based document. "
            "Please upload a text-based PDF or paste the notes manually."
        )

    return text


# ============================================================
# DOCX
# ============================================================

def extract_docx_text(content):
    try:
        document = Document(BytesIO(content))
    except Exception as exc:
        raise UnsupportedFileError(
            "The DOCX file could not be opened."
        ) from exc

    parts = []

    # Normal paragraphs
    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Tables
    for table in document.tables:

        for row in table.rows:

            cells = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:
                    cells.append(text)

            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts).strip()

    if not text:
        raise UnsupportedFileError(
            "This DOCX file does not contain readable text."
        )

    return text


# ============================================================
# HTML
# ============================================================

def extract_html_text(content):
    try:
        decoded = decode_text(content)
    except UnsupportedFileError:
        raise UnsupportedFileError(
            "Could not read the HTML file."
        )

    soup = BeautifulSoup(
        decoded,
        "html.parser",
    )

    # Remove unnecessary elements.
    for element in soup(
        ["script", "style", "noscript"]
    ):
        element.decompose()

    text = soup.get_text(
        "\n",
        strip=True,
    )

    if not text:
        raise UnsupportedFileError(
            "The HTML file does not contain readable text."
        )

    return text


# ============================================================
# PLAIN TEXT / CSV / JSON / MARKDOWN / CODE
# ============================================================

def extract_plain_text(content):
    text = decode_text(content)

    if not text.strip():
        raise UnsupportedFileError(
            "The uploaded file does not contain readable text."
        )

    return text.strip()


def decode_text(content):
    """
    Try common encodings.
    """

    for encoding in (
        "utf-8",
        "utf-8-sig",
        "utf-16",
        "utf-16-le",
        "utf-16-be",
        "latin-1",
    ):

        try:

            text = content.decode(
                encoding
            ).strip()

        except UnicodeDecodeError:
            continue

        if text and has_enough_text(text):
            return text

    raise UnsupportedFileError(
        "Could not read text from this file. "
        "Please upload PDF, DOCX, TXT, CSV, HTML, Markdown, "
        "or paste the notes manually."
    )


def has_enough_text(text):
    """
    Make sure the decoded content actually looks like text.
    """

    if not text:
        return False

    printable = sum(
        1
        for char in text
        if char.isprintable() or char.isspace()
    )

    ratio = printable / max(
        len(text),
        1,
    )

    return ratio > 0.85